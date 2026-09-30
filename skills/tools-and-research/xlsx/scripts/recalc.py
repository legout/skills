# /// script
# dependencies = ["openpyxl", "defusedxml"]
# ///
"""Recalculate an XLSX copy with isolated LibreOffice, checking cached results.

Refuses external workbook links, macros and unsupported spilling formulas.
On any error the requested output is not published; JSON + nonzero exit is returned.
"""
import argparse
import json
import posixpath
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from defusedxml import ElementTree as ET
from office import convert, new_output, validate, REL

SPILL = re.compile(r"\b(?:_xlfn\.)?(?:_xlws\.)?(FILTER|SORT|SORTBY|UNIQUE|SEQUENCE|RANDARRAY|XLOOKUP|XMATCH)\s*\(", re.I)


def recalculate(source, output, timeout=120):
    source, output = new_output(source, output)
    if source.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
        raise ValueError("Recalculation supports .xlsx only; use Excel for macro/template files")
    structural = validate(source)
    if structural["errors"]:
        raise ValueError(structural["errors"])
    with zipfile.ZipFile(source) as archive:
        if any(name.startswith("xl/externalLinks/") for name in archive.namelist()):
            raise ValueError("External workbook links detected; recalculate in Excel with linked files available")
    original = load_workbook(source, data_only=False)
    formulas = {(sheet.title, cell.coordinate): cell.value for sheet in original for row in sheet for cell in row if cell.data_type == "f"}
    if any(not isinstance(value, str) or SPILL.search(value) for value in formulas.values()):
        raise ValueError("Array/spilling or engine-sensitive lookup formula detected; verify in Excel instead")
    original.close()
    with tempfile.TemporaryDirectory(prefix="xlsx-recalc-") as temporary:
        work = Path(temporary)
        input_dir, output_dir = work / "input", work / "output"
        input_dir.mkdir()
        output_dir.mkdir()
        copied = input_dir / source.name
        shutil.copy2(source, copied)
        calculated = convert(copied, output_dir, "xlsx", timeout)
        structural = validate(calculated)
        if structural["errors"]:
            raise ValueError(structural["errors"])
        strings = load_workbook(calculated, data_only=False)
        cached = load_workbook(calculated, data_only=True)
        # Parse blank formula caches once per sheet, not once per empty cell.
        empty_text = set()
        with zipfile.ZipFile(calculated) as package:
            ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            workbook = ET.fromstring(package.read("xl/workbook.xml"))
            rels = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
            targets = {e.get("Id"): e.get("Target") for e in rels.findall(f"{{{REL}}}Relationship")}
            for sheet in workbook.findall("s:sheets/s:sheet", ns):
                target = targets.get(sheet.get(f"{{{rel_ns}}}id"))
                if not target:
                    raise ValueError("Missing worksheet relationship target")
                member = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
                tree = ET.fromstring(package.read(member))
                for element in tree.findall(".//s:c", ns):
                    if element.get("t") in {"str", "inlineStr"}:
                        coordinate = (sheet.get("name"), element.get("r"))
                        if coordinate in formulas:
                            empty_text.add(coordinate)
        errors = []
        for sheet in cached:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "e":
                        errors.append({"sheet": sheet.title, "cell": cell.coordinate, "error": cell.value})
        for (sheet, coordinate), formula in formulas.items():
            if sheet not in strings.sheetnames or strings[sheet][coordinate].data_type != "f":
                errors.append({"sheet": sheet, "cell": coordinate, "error": "Formula lost during conversion"})
            elif cached[sheet][coordinate].value is None:
                # LibreOffice caches ="" as empty text, which openpyxl exposes as None.
                if (sheet, coordinate) not in empty_text:
                    errors.append({"sheet": sheet, "cell": coordinate, "error": "Missing calculated value"})
        strings.close()
        cached.close()
        if errors:
            return {"status": "errors", "formulas": len(formulas), "errors": errors, "output": None}
        with output.open("xb") as stream, calculated.open("rb") as origin:
            shutil.copyfileobj(origin, stream)
    return {"status": "ok", "formulas": len(formulas), "errors": [], "output": str(output), "engine": "LibreOffice", "note": "Evaluation checked, not business correctness or Excel layout fidelity"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--timeout", type=int, default=120)
    args = parser.parse_args()
    try:
        report = recalculate(args.input, args.output, args.timeout)
        print(json.dumps(report))
        return 1 if report["errors"] else 0
    except Exception as exc:
        print(json.dumps({"status": "error", "error": str(exc)}))
        return 2


if __name__ == "__main__":
    sys.exit(main())
