# /// script
# dependencies = ["python-docx>=1.2", "python-pptx", "openpyxl", "pypdf[crypto]>=6.19", "reportlab", "defusedxml", "pypdfium2", "pillow"]
# ///
"""Real-file regression checks for the four original Office skill tools."""
import json
import subprocess
import sys
import tempfile
import shutil
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1] / "skills/tools-and-research"


def run(skill, script, *args, ok=True):
    result = subprocess.run([sys.executable, str(ROOT / skill / "scripts" / script), *map(str, args)], capture_output=True, text=True)
    assert (result.returncode == 0) == ok, (result.returncode, result.stdout, result.stderr)
    return json.loads(result.stdout)


def main():
    with tempfile.TemporaryDirectory() as temporary:
        d = Path(temporary)
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Old ").bold = True
        p.add_run("word and tail")
        doc.save(d / "original.docx")
        assert run("docx", "office.py", "validate", d / "original.docx")["status"] == "ok"
        run("docx", "edit.py", "replace", d / "original.docx", d / "tracked.docx", "Old word", "New phrase", "--author", "Tester")
        run("docx", "edit.py", "accept", d / "tracked.docx", d / "accepted.docx")
        assert Document(d / "accepted.docx").paragraphs[0].text == "New phrase and tail"
        run("docx", "edit.py", "comment", d / "accepted.docx", d / "commented.docx", "New phrase", "Check this", "--author", "Tester")
        assert len(Document(d / "commented.docx").comments) == 1
        run("docx", "edit.py", "accept", d / "tracked.docx", d / "original.docx", ok=False)
        # Accepting a deleted paragraph mark must join, not leave an empty bullet.
        joined = Document()
        first = joined.add_paragraph("First ")
        joined.add_paragraph("second", style="Title")
        props = OxmlElement("w:rPr")
        props.append(OxmlElement("w:del"))
        first._p.get_or_add_pPr().append(props)
        joined.save(d / "join.docx")
        run("docx", "edit.py", "accept", d / "join.docx", d / "joined.docx")
        assert [p.text for p in Document(d / "joined.docx").paragraphs] == ["First second"]
        assert Document(d / "joined.docx").paragraphs[0].style.name == "Title"
        # Unresolved relationships and unsafe XML must fail, not just reopen.
        with zipfile.ZipFile(d / "original.docx") as source, zipfile.ZipFile(d / "broken.docx", "w") as broken:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == "_rels/.rels":
                    content = content.replace(b"word/document.xml", b"word/missing.xml")
                broken.writestr(info, content)
        assert run("docx", "office.py", "validate", d / "broken.docx", ok=False)["errors"]
        with zipfile.ZipFile(d / "original.docx") as source, zipfile.ZipFile(d / "entity.docx", "w") as unsafe:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == "word/document.xml":
                    content = b'<!DOCTYPE document [<!ENTITY secret SYSTEM "file:///etc/passwd">]><document>&secret;</document>'
                unsafe.writestr(info, content)
        run("docx", "office.py", "validate", d / "entity.docx", ok=False)
        run("docx", "office.py", "template", d / "original.docx", d / "template.dotx")
        run("docx", "office.py", "template", d / "template.dotx", d / "from-template.docx")
        assert Document(d / "from-template.docx").paragraphs[0].text == "Old word and tail"

        slides = Presentation()
        page = slides.slides.add_slide(slides.slide_layouts[6])
        page.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Slide one"
        data = CategoryChartData()
        data.categories = ["A", "B"]
        data.add_series("Series", [1, 2])
        page.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), data)
        page.notes_slide.notes_text_frame.text = "Speaker note"
        hidden = slides.slides.add_slide(slides.slide_layouts[6])
        hidden._element.set("show", "0")
        slides.save(d / "slides.pptx")
        run("pptx", "slides.py", d / "slides.pptx", d / "reordered.pptx", "--order", "2,1,1")
        deck = Presentation(d / "reordered.pptx")
        assert len(deck.slides) == 3 and deck.slides[1].shapes[0].text == "Slide one"
        deck.slides[1].shapes[0].text = "Independent"
        assert deck.slides[2].shapes[0].text == "Slide one"
        changed_data = CategoryChartData()
        changed_data.categories = ["A", "B"]
        changed_data.add_series("Series", [8, 9])
        deck.slides[1].shapes[1].chart.replace_data(changed_data)
        assert list(deck.slides[2].shapes[1].chart.series[0].values) == [1, 2]
        assert deck.slides[2].notes_slide.notes_text_frame.text == "Speaker note"
        run("pptx", "office.py", "validate", d / "reordered.pptx")
        run("pptx", "office.py", "template", d / "slides.pptx", d / "deck-template.potx")
        run("pptx", "office.py", "template", d / "deck-template.potx", d / "from-template.pptx")
        assert len(Presentation(d / "from-template.pptx").slides) == 2

        book = Workbook()
        book.active.append(["Value", "Twice"])
        book.active.append([3, "=A2*2"])
        book.save(d / "formulas.xlsx")
        run("xlsx", "office.py", "template", d / "formulas.xlsx", d / "sheet-template.xltx")
        assert load_workbook(d / "sheet-template.xltx").template
        result = run("xlsx", "recalc.py", d / "formulas.xlsx", d / "calculated.xlsx")
        assert result["status"] == "ok" and load_workbook(d / "calculated.xlsx", data_only=True).active["B2"].value == 6
        book.active["B2"] = "=1/0"
        book.save(d / "bad.xlsx")
        assert run("xlsx", "recalc.py", d / "bad.xlsx", d / "bad-output.xlsx", ok=False)["errors"]
        assert not (d / "bad-output.xlsx").exists()
        book.active["B2"] = '=IF(A2=3,"",1)'
        book.save(d / "empty.xlsx")
        run("xlsx", "recalc.py", d / "empty.xlsx", d / "empty-result.xlsx")
        book.active["B2"] = "=SEQUENCE(3)"
        book.save(d / "spill.xlsx")
        run("xlsx", "recalc.py", d / "spill.xlsx", d / "spill-result.xlsx", ok=False)

        pdf = canvas.Canvas(str(d / "form.pdf"))
        pdf.drawString(40, 740, "Form")
        from PIL import Image
        Image.new("RGB", (10, 10), "red").save(d / "logo.png")
        pdf.drawImage(str(d / "logo.png"), 350, 650, width=30, height=30)
        pdf.acroForm.textfield(name="name", x=40, y=650, width=200, height=25, maxlen=12)
        pdf.acroForm.checkbox(name="approved", x=40, y=600, checked=False)
        pdf.acroForm.choice(name="period", x=40, y=500, width=200, height=25, options=["Q1", "Q2"], value="Q1")
        pdf.acroForm.radio(name="priority", value="low", x=40, y=460, selected=True)
        pdf.acroForm.radio(name="priority", value="high", x=80, y=460, selected=False)
        pdf.showPage()
        pdf.save()
        fields = run("pdf", "pdf.py", "fields", d / "form.pdf")
        assert "name" in fields["fields"]
        (d / "values.json").write_text(json.dumps({"name": "Ada", "approved": "/Yes", "period": "Q2", "priority": "/high"}))
        run("pdf", "pdf.py", "fill", d / "form.pdf", d / "filled.pdf", "--values", d / "values.json")
        assert PdfReader(d / "filled.pdf").get_fields()["name"]["/V"] == "Ada"
        assert PdfReader(d / "filled.pdf").get_fields()["approved"]["/V"] == "/Yes"
        assert PdfReader(d / "filled.pdf").get_fields()["period"]["/V"] == "Q2"
        assert PdfReader(d / "filled.pdf").get_fields()["priority"]["/V"] == "/high"
        run("pdf", "pdf.py", "flatten", d / "filled.pdf", d / "flat.pdf")
        assert not PdfReader(d / "flat.pdf").get_fields()
        assert "Ada" in PdfReader(d / "flat.pdf").pages[0].extract_text()
        assert run("pdf", "pdf.py", "images", d / "filled.pdf", d / "embedded-images")["images"] == 1
        mark = canvas.Canvas(str(d / "watermark.pdf"))
        mark.drawString(350, 740, "DRAFT")
        mark.save()
        run("pdf", "pdf.py", "watermark", d / "filled.pdf", d / "marked.pdf", "--overlay", d / "watermark.pdf")
        assert "DRAFT" in PdfReader(d / "marked.pdf").pages[0].extract_text()
        run("pdf", "pdf.py", "select", d / "watermark.pdf", d / "double.pdf", "--pages", "1,1")
        assert len(PdfReader(d / "double.pdf").pages) == 2
        (d / "long.json").write_text(json.dumps({"name": "x" * 20}))
        run("pdf", "pdf.py", "fill", d / "form.pdf", d / "overlong.pdf", "--values", d / "long.json", ok=False)
        (d / "values.json").write_text(json.dumps({"missing": "value"}))
        run("pdf", "pdf.py", "fill", d / "form.pdf", d / "unknown.pdf", "--values", d / "values.json", ok=False)
        assert not (d / "unknown.pdf").exists()
        (d / "annotations.json").write_text(json.dumps([{"page": 1, "x": 40, "y": 550, "text": "Reviewed"}]))
        run("pdf", "pdf.py", "stamp", d / "filled.pdf", d / "stamped.pdf", "--annotations", d / "annotations.json")
        assert "Reviewed" in PdfReader(d / "stamped.pdf").pages[0].extract_text()
        (d / "password.txt").write_text("test-password")
        run("pdf", "pdf.py", "encrypt", d / "filled.pdf", d / "secret.pdf", "--new-password-file", d / "password.txt")
        assert PdfReader(d / "secret.pdf").is_encrypted
        run("pdf", "pdf.py", "decrypt", d / "secret.pdf", d / "plain.pdf", "--password-file", d / "password.txt")
        assert not PdfReader(d / "plain.pdf").is_encrypted
        run("pdf", "pdf.py", "render", d / "filled.pdf", d / "pdf-images")
        run("pdf", "pdf.py", "render", d / "flat.pdf", d / "flat-images")
        assert (d / "pdf-images/page-001.png").exists()
        picture = Image.open(d / "pdf-images/page-001.png").convert("RGB")
        assert any(blue > red + 10 for count, (red, green, blue) in picture.getcolors(picture.width * picture.height)), "Form appearances were not rendered"
        flat_picture = Image.open(d / "flat-images/page-001.png").convert("RGB")
        radios = [item for item in fields["widgets"] if item["name"] == "priority"]
        def dark_center(image, item):
            left, bottom, right, top = item["rect"]
            x = (left + right) / 2 * 120 / 72
            y = image.height - (bottom + top) / 2 * 120 / 72
            center = image.crop((int(x - 4), int(y - 4), int(x + 4), int(y + 4)))
            return sum(count for count, rgb in center.getcolors(64) if max(rgb) < 100)
        assert dark_center(picture, radios[1]) > dark_center(picture, radios[0]) + 5
        assert dark_center(flat_picture, radios[1]) > dark_center(flat_picture, radios[0]) + 5, "Flattening lost radio selection"

        for skill, source in (("docx", "commented.docx"), ("pptx", "reordered.pptx"), ("xlsx", "calculated.xlsx")):
            report = run(skill, "office.py", "render", d / source, d / f"{skill}-images")
            assert report["pages"] > 0 and (d / f"{skill}-images/contact-sheet.png").exists()
            if skill == "pptx":
                assert report["pages"] == 3, "Hidden slides omitted from preview"
        # Standalone copies of the shared helper must not drift.
        common = (ROOT / "docx/scripts/office.py").read_bytes()
        assert all((ROOT / f"{skill}/scripts/office.py").read_bytes() == common for skill in ("pptx", "xlsx", "pdf"))
        if len(sys.argv) == 2:
            shutil.copytree(d, Path(sys.argv[1]))
        print("Office skills: real DOCX tracked edits/comments, PPTX ordering/independent chart copies, XLSX recalculation/errors, PDF forms/encryption/rendering and Office renders passed.")


if __name__ == "__main__":
    main()
